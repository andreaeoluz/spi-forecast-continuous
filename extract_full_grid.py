#!/usr/bin/env python3
"""Dump the full (p,q) validation grid (WI, RMSE, MAE) for every region and
strategy from outputs/<region>/grid_search/results/spi3_tl_<flag>/
grid_search_results_spi3_tl_<flag>.xlsx into one tidy CSV."""
import openpyxl
import csv
from pathlib import Path

REGIONS = [
    ("Norte", "North"),
    ("Nordeste", "Northeast"),
    ("Centro-Oeste", "Center-West"),
    ("Sudeste", "Southeast"),
    ("Sul", "South"),
]
STRATS = [("True", "TL"), ("False", "Scratch")]

OUT = Path("outputs")
rows_out = []

for region_pt, region_en in REGIONS:
    for tl_flag, strat_label in STRATS:
        p = OUT / region_pt / "grid_search" / "results" / f"spi3_tl_{tl_flag}" / f"grid_search_results_spi3_tl_{tl_flag}.xlsx"
        if not p.exists():
            print("MISSING", p)
            continue
        wb = openpyxl.load_workbook(p, data_only=True)
        ws = wb.active
        header = [c.value for c in ws[1]]
        idx = {h: i for i, h in enumerate(header)}
        for row in ws.iter_rows(min_row=2, values_only=True):
            rows_out.append({
                "region": region_en, "strategy": strat_label,
                "p": row[idx["p"]], "q": row[idx["q"]],
                "wi": row[idx["best_wi"]], "rmse": row[idx["best_rmse"]],
                "mae": row[idx["best_mae"]],
            })

with open("grid_full.csv", "w", newline="") as f:
    w = csv.DictWriter(f, fieldnames=["region", "strategy", "p", "q", "wi", "rmse", "mae"])
    w.writeheader()
    w.writerows(rows_out)

print(f"wrote {len(rows_out)} rows to grid_full.csv")
